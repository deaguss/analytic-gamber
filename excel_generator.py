import json
import pandas as pd
from pathlib import Path
from datetime import datetime

CAT_FILL = {
    'JUDI':       'FFCCCC',   # merah muda
    'TIDAK_JUDI': 'CCFFCC',   # hijau muda
    'AMBIGU':     'FFF2CC',   # kuning muda
    'RANDOM':     'FFD580',   # oranye muda
    'UNKNOWN':    'F2F2F2',   # abu-abu
}
CAT_LABEL = {
    'JUDI':       'Berindikasi Judi',
    'TIDAK_JUDI': 'Tidak Judi',
    'AMBIGU':     'Ambigu',
    'RANDOM':     'Data Random',
    'UNKNOWN':    'Tidak Diketahui',
}

# Legend inline di bawah data
LEGEND_ROWS = [
    ('JUDI',       'Berindikasi Judi', 'Merah muda  — video berindikasi komentar judi online dominan'),
    ('TIDAK_JUDI', 'Tidak Judi',       'Hijau muda  — video dengan komentar dominan bukan judi'),
    ('AMBIGU',     'Ambigu',           'Kuning muda — video dengan komentar campuran / tidak jelas'),
    ('RANDOM',     'Data Random',      'Oranye muda — video acak untuk uji efektivitas deteksi cascading'),
]


def apply_colors_with_legend(writer, df: pd.DataFrame, sheet_name: str):
    """Warnai baris data berdasarkan kategori & tambahkan legenda di bawah data."""
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter

    code_series = df['_cat_code'].tolist() if '_cat_code' in df.columns else []
    display     = df.drop(columns=['_cat_code'], errors='ignore')

    display.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]

    # Bold + center header
    for cell in ws[1]:
        cell.font      = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Auto-width kolom
    for ci, col in enumerate(display.columns, 1):
        vals    = display[col].astype(str).tolist() + [str(col)]
        max_len = max(len(v) for v in vals) + 4
        ws.column_dimensions[get_column_letter(ci)].width = min(max_len, 45)

    # Warnai baris data
    for ri, cat in enumerate(code_series, start=2):
        fill = PatternFill(
            start_color=CAT_FILL.get(cat, 'F2F2F2'),
            end_color=CAT_FILL.get(cat, 'F2F2F2'),
            fill_type='solid'
        )
        for ci in range(1, len(display.columns) + 1):
            ws.cell(row=ri, column=ci).fill = fill

    # Legenda di bawah data
    last_data_row = len(code_series) + 1  # +1 untuk header
    legend_start  = last_data_row + 2     # 1 baris kosong pemisah

    # Judul legenda
    title_cell = ws.cell(row=legend_start, column=1, value="LEGENDA WARNA KATEGORI")
    title_cell.font = Font(bold=True, size=11)

    # Header kolom legenda
    legend_header_row = legend_start + 1
    for ci, h in enumerate(['Warna', 'Kode', 'Label', 'Keterangan'], 1):
        c = ws.cell(row=legend_header_row, column=ci, value=h)
        c.font      = Font(bold=True)
        c.alignment = Alignment(horizontal='center')

    # Isi legenda
    for i, (code, label, desc) in enumerate(LEGEND_ROWS, start=legend_header_row + 1):
        fill = PatternFill(
            start_color=CAT_FILL.get(code, 'F2F2F2'),
            end_color=CAT_FILL.get(code, 'F2F2F2'),
            fill_type='solid'
        )
        # Kolom 1: kotak warna (kosong, hanya warna)
        ws.cell(row=i, column=1, value="").fill = fill
        # Kolom 2: kode
        c2 = ws.cell(row=i, column=2, value=code)
        c2.fill = fill
        # Kolom 3: label
        c3 = ws.cell(row=i, column=3, value=label)
        c3.fill = fill
        # Kolom 4: keterangan
        c4 = ws.cell(row=i, column=4, value=desc)
        c4.fill = fill

    # Lebar kolom legenda
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 65


def generate_excel():
    results_dir = Path("results")
    json_path   = results_dir / "aggregate_results.json"

    if not json_path.exists():
        print(f"[ERROR] File {json_path} tidak ditemukan.")
        print("Jalankan detector.py terlebih dahulu.")
        return

    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    videos = data.get('per_video_results', [])
    if not videos:
        print("Tidak ada data video untuk diekspor.")
        return

    print(f"Loaded {len(videos)} video dari aggregate_results.json")

    def code(v):  return v.get('sample_category', 'UNKNOWN')
    def label(v): return CAT_LABEL.get(code(v), code(v))

    # 1. Rekap Keseluruhan
    df_rekap = pd.DataFrame(videos)
    if 'sample_category' not in df_rekap.columns:
        df_rekap['sample_category'] = 'UNKNOWN'
    df_rekap['Kategori Sampel'] = df_rekap['sample_category'].map(lambda c: CAT_LABEL.get(c, c))
    df_rekap['_cat_code']       = df_rekap['sample_category']
    cols = ['Kategori Sampel', '_cat_code'] + [
        c for c in df_rekap.columns if c not in ('Kategori Sampel', '_cat_code', 'sample_category')
    ]
    df_rekap = df_rekap[cols]

    # 2. Agreement
    df_agreement = pd.DataFrame([{
        'Kategori Sampel'        : label(v),
        'Video ID'               : v.get('video_id'),
        'Komentar Ditarik'       : v.get('total_comments'),
        'Anomali (Masuk AI)'     : v.get('rule_based_ambigu'),
        'Gemini (Judi)'          : v.get('gemini_judi'),
        'Gemini Sepakat'         : v.get('gemini_agree_count'),
        'Gemini Tidak Sepakat'   : v.get('gemini_disagree_count'),
        'Gemini Agreement Rate %': v.get('gemini_agreement_rate'),
        'GPT (Judi)'             : v.get('gpt_judi'),
        'GPT Sepakat'            : v.get('gpt_agree_count'),
        'GPT Tidak Sepakat'      : v.get('gpt_disagree_count'),
        'GPT Agreement Rate %'   : v.get('gpt_agreement_rate'),
        '_cat_code'              : code(v),
    } for v in videos])

    # 3. Confidence
    df_confidence = pd.DataFrame([{
        'Kategori Sampel'      : label(v),
        'Video ID'             : v.get('video_id'),
        'Anomali (Masuk AI)'   : v.get('rule_based_ambigu'),
        'Gemini Avg Confidence': v.get('gemini_avg_confidence'),
        'GPT Avg Confidence'   : v.get('gpt_avg_confidence'),
        '_cat_code'            : code(v),
    } for v in videos])

    # 4. Latency
    df_latency = pd.DataFrame([{
        'Kategori Sampel'          : label(v),
        'Video ID'                 : v.get('video_id'),
        'Anomali (Masuk AI)'       : v.get('rule_based_ambigu'),
        'Gemini Avg Latency (ms)'  : v.get('gemini_avg_latency'),
        'GPT Avg Latency (ms)'     : v.get('gpt_avg_latency'),
        'Total Processing Time (s)': v.get('processing_time_seconds'),
        '_cat_code'                : code(v),
    } for v in videos])

    # 5. Cost
    df_cost = pd.DataFrame([{
        'Kategori Sampel'     : label(v),
        'Video ID'            : v.get('video_id'),
        'Anomali (Masuk AI)'  : v.get('rule_based_ambigu'),
        'Gemini Input Tokens' : v.get('gemini_total_input_tokens', 0),
        'Gemini Output Tokens': v.get('gemini_total_output_tokens', 0),
        'Gemini Cost (USD)'   : v.get('gemini_total_cost_usd', 0.0),
        'GPT Input Tokens'    : v.get('gpt_total_input_tokens', 0),
        'GPT Output Tokens'   : v.get('gpt_total_output_tokens', 0),
        'GPT Cost (USD)'      : v.get('gpt_total_cost_usd', 0.0),
        '_cat_code'           : code(v),
    } for v in videos])

    file_map = [
        (df_rekap,      "1_Rekap_Keseluruhan.xlsx"),
        (df_agreement,  "2_Report_Agreement.xlsx"),
        (df_confidence, "3_Report_Confidence.xlsx"),
        (df_latency,    "4_Report_Latency.xlsx"),
        (df_cost,       "5_Report_Cost.xlsx"),
    ]

    for df, fname in file_map:
        path = results_dir / fname
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            apply_colors_with_legend(writer, df, 'Data')
        print(f"  Saved: {fname}")

    print(f"\nBerhasil membuat 5 file Excel  ->  {results_dir}/")
    print("  Warna baris:")
    print("    Merah muda  = JUDI")
    print("    Hijau muda  = TIDAK_JUDI")
    print("    Kuning muda = AMBIGU")
    print("    Oranye muda = RANDOM")
    print("  Legenda warna tampil di bawah data pada setiap sheet.")


if __name__ == "__main__":
    generate_excel()